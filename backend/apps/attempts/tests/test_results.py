"""Staff results API: scoping (branch + teacher subject), filters, detail, summary, export."""

import io
from datetime import timedelta

from django.utils import timezone
from openpyxl import load_workbook

from apps.attempts.models import Attempt
from apps.attempts.scoping import attempts_for
from apps.attempts.services import results
from apps.content.models import Test

from .helpers import PublicAPITestCase, make_published

R = "/api/v1/results"


class ResultsBase(PublicAPITestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.test_a1 = make_published(cls.subj_a1, "Algebra", pass_percent=50)
        cls.test_a2 = make_published(cls.subj_a2, "Mexanika")
        cls.test_b = make_published(cls.subj_b1, "B testi")

    def setUp(self):
        super().setUp()
        self.att_a1 = self.done(self.test_a1, "Ali Valiyev", right=True)
        self.att_a1_bad = self.done(self.test_a1, "Vali Aliyev", right=False)
        self.att_a2 = self.done(self.test_a2, "Fizik Bola")
        self.att_b = self.done(self.test_b, "Boshqa Filial")

    def done(self, test, name, right=True):
        """Run one complete anonymous attempt through the public API (its own device)."""
        from rest_framework.test import APIClient

        client = APIClient()
        body = self.start(test, client=client, name=name).json()
        token = body["access_token"]
        if right:
            for item in body["items"]:
                q = test.questions.get(order=item["order"] - 1)
                self.answer(token, item, self.correct_ids(q), client=client)
        else:
            for item in body["items"]:
                q = test.questions.get(order=item["order"] - 1)
                wrong = [str(o.id) for o in q.options.filter(is_correct=False)][:1]
                self.answer(token, item, wrong, client=client)
        client.post(f"/api/v1/public/attempts/{token}/finish/")
        return Attempt.objects.get(access_token=token)

    def ids(self, response):
        return {row["id"] for row in response.json()["results"]}


class ScopingTests(ResultsBase):
    def test_superadmin_sees_all(self):
        r = self.as_user(self.root).get(f"{R}/")
        self.assertEqual(r.json()["count"], 4)

    def test_admin_sees_only_own_branch(self):
        r = self.as_user(self.admin_a).get(f"{R}/")
        self.assertEqual(self.ids(r), {str(self.att_a1.id), str(self.att_a1_bad.id), str(self.att_a2.id)})
        self.assertNotIn(str(self.att_b.id), self.ids(r))

    def test_admin_cannot_open_other_branch_result(self):
        c = self.as_user(self.admin_a)
        self.assertEqual(c.get(f"{R}/{self.att_b.id}/").status_code, 404)
        self.assertEqual(self.as_user(self.admin_b).get(f"{R}/{self.att_b.id}/").status_code, 200)

    def test_admin_branch_filter_cannot_widen_scope(self):
        r = self.as_user(self.admin_a).get(f"{R}/", {"branch": str(self.b.id)})
        self.assertEqual(r.json()["count"], 0)

    def test_teacher_only_assigned_subjects(self):
        r = self.as_user(self.teacher_a).get(f"{R}/")
        self.assertEqual(self.ids(r), {str(self.att_a1.id), str(self.att_a1_bad.id)})
        self.assertEqual(self.as_user(self.teacher_a).get(f"{R}/{self.att_a2.id}/").status_code, 404)  # subject not assigned
        self.assertEqual(self.as_user(self.teacher_a).get(f"{R}/{self.att_b.id}/").status_code, 404)

    def test_teacher_of_other_branch_sees_nothing_of_a(self):
        r = self.as_user(self.teacher_b).get(f"{R}/")
        self.assertEqual(self.ids(r), {str(self.att_b.id)})

    def test_summary_and_export_are_scoped(self):
        t = self.as_user(self.teacher_a)
        self.assertEqual(t.get(f"{R}/tests/{self.test_a2.id}/summary/").status_code, 404)
        self.assertEqual(self.as_user(self.admin_a).get(f"{R}/tests/{self.test_b.id}/summary/").status_code, 404)
        wb = load_workbook(io.BytesIO(self.as_user(self.admin_a).get(f"{R}/export/").content))
        names = {row[3] for row in wb.active.iter_rows(min_row=2, values_only=True)}
        self.assertNotIn("Boshqa Filial", names)
        self.assertEqual(len(names), 3)

    def test_anonymous_refused(self):
        self.assertIn(self.anon.get(f"{R}/").status_code, (401, 403))
        self.assertIn(self.anon.get(f"{R}/export/").status_code, (401, 403))

    def test_helper_none_for_anonymous_user(self):
        from django.contrib.auth.models import AnonymousUser

        self.assertEqual(attempts_for(AnonymousUser()).count(), 0)


class ListFilterTests(ResultsBase):
    def get(self, **params):
        return self.as_user(self.root).get(f"{R}/", params)

    def test_list_shape_and_summary(self):
        body = self.get(test=str(self.test_a1.id)).json()
        self.assertEqual(body["count"], 2)
        self.assertEqual(body["summary"]["attempts"], 2)
        self.assertEqual(body["summary"]["pass_rate"], 50.0)
        row = body["results"][0]
        for key in ("full_name", "age", "test", "subject", "branch", "percent", "passed", "duration_sec"):
            self.assertIn(key, row)
        self.assertNotIn("access_token", row)

    def test_filters(self):
        self.assertEqual(self.get(passed="true").json()["count"], 3)
        self.assertEqual(self.get(passed="false").json()["count"], 1)
        self.assertEqual(self.get(subject=str(self.subj_a2.id)).json()["count"], 1)
        self.assertEqual(self.get(branch=str(self.b.id)).json()["count"], 1)
        self.assertEqual(self.get(search="vali").json()["count"], 2)
        self.assertEqual(self.get(status="expired").json()["count"], 0)

    def test_date_range(self):
        Attempt.objects.filter(pk=self.att_a1.pk).update(started_at=timezone.now() - timedelta(days=10))
        today = timezone.localdate().isoformat()
        self.assertEqual(self.get(date_from=today).json()["count"], 3)
        old = (timezone.localdate() - timedelta(days=5)).isoformat()
        self.assertEqual(self.get(date_to=old).json()["count"], 1)

    def test_invalid_filter_is_400_not_500(self):
        self.assertEqual(self.get(test="nope").status_code, 400)
        self.assertEqual(self.get(date_from="32-13-2020").status_code, 400)

    def test_newest_first_and_paginated(self):
        body = self.get(page_size=2).json()
        self.assertEqual(len(body["results"]), 2)
        self.assertIsNotNone(body["next"])
        times = [r["started_at"] for r in self.get().json()["results"]]
        self.assertEqual(times, sorted(times, reverse=True))

    def test_in_progress_hidden_by_default(self):
        self.start(self.test_a1, name="Yechayotgan")
        self.assertEqual(self.get().json()["count"], 4)
        self.assertEqual(self.get(status="in_progress").json()["count"], 1)


class DetailAndSummaryTests(ResultsBase):
    def test_detail_uses_snapshot_after_test_edit(self):
        # Edit the live question afterwards: the stored result must not change.
        q = self.test_a1.questions.first()
        q.options.update(is_correct=False)
        q.body_html = "<p>OZGARGAN</p>"
        q.save()
        d = self.as_user(self.admin_a).get(f"{R}/{self.att_a1.id}/").json()
        self.assertEqual(d["full_name"], "Ali Valiyev")
        self.assertEqual(len(d["items"]), 2)
        first = next(i for i in d["items"] if i["type"] == "single")
        self.assertTrue(first["is_correct"])
        self.assertTrue(any(o["is_correct"] and o["selected"] for o in first["options"]))
        self.assertNotIn("OZGARGAN", str(d))

    def test_detail_wrong_answer_flags(self):
        d = self.as_user(self.admin_a).get(f"{R}/{self.att_a1_bad.id}/").json()
        self.assertFalse(d["passed"])
        self.assertTrue(all(i["is_correct"] is False for i in d["items"]))

    def test_test_summary(self):
        body = self.as_user(self.admin_a).get(f"{R}/tests/{self.test_a1.id}/summary/").json()
        self.assertEqual(body["attempts"], 2)
        self.assertEqual(body["pass_rate"], 50.0)
        self.assertEqual(body["avg_percent"], 50.0)
        q1, q2 = body["questions"]
        self.assertEqual((q1["answered_count"], q1["correct_count"], q1["correct_percent"]), (2, 1, 50.0))
        self.assertIsNotNone(q1["top_wrong_option"])
        self.assertEqual(q1["order"], 1)

    def test_empty_summary(self):
        test = Test.objects.create(subject=self.subj_a1, title="Bo'sh", status="draft")
        body = self.as_user(self.admin_a).get(f"{R}/tests/{test.id}/summary/").json()
        self.assertEqual(body["attempts"], 0)
        self.assertIsNone(body["avg_percent"])
        self.assertEqual(body["questions"], [])


class ExportTests(ResultsBase):
    def test_xlsx_columns_and_values(self):
        r = self.as_user(self.root).get(f"{R}/export/", {"test": str(self.test_a1.id)})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:5], ("Filial", "Fan", "Test", "Ism", "Yosh"))
        self.assertEqual(len(rows), 3)
        self.assertEqual(len(rows[0]), 13)

    def test_formula_injection_neutralised(self):
        Attempt.objects.filter(pk=self.att_a1.pk).update(full_name="=HYPERLINK(\"http://x\")")
        Test.objects.filter(pk=self.test_a1.pk).update(title="@SUM(1)")
        r = self.as_user(self.root).get(f"{R}/export/", {"test": str(self.test_a1.id)})
        ws = load_workbook(io.BytesIO(r.content)).active
        cells = [c for row in ws.iter_rows(min_row=2) for c in row]
        for cell in cells:
            if isinstance(cell.value, str):
                self.assertNotIn(cell.value[:1], "=+-@")
                self.assertEqual(cell.data_type, "s")
        self.assertTrue(any(str(c.value).startswith("'=HYPERLINK") for c in cells))

    def test_row_limit(self):
        from unittest import mock

        with mock.patch.object(results, "EXPORT_LIMIT", 2):
            r = self.as_user(self.root).get(f"{R}/export/")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Filtrlarni toraytiring", r.json()["detail"])
